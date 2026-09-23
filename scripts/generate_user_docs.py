#!/usr/bin/env python3
"""Regenerate shipped guides: pip install -r scripts/docs-requirements.txt."""
from pathlib import Path
from html import escape
from urllib.parse import urlsplit
import shutil
import markdown
from lxml import html
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / 'frontend/public'
SOURCE_URL = 'https://github.com/waptug/ScriptStudio/blob/master/'
CSS = '''body{font-family:system-ui,sans-serif;color:#203040;line-height:1.6;max-width:920px;margin:auto;padding:32px 24px;background:white}h1,h2,h3{color:#164e63;line-height:1.25}h2{padding-top:20px;border-bottom:1px solid #c3d0d8}a{color:#155cb0}pre,code{background:#edf2f5;overflow-wrap:anywhere}pre{padding:16px;white-space:pre-wrap}table{border-collapse:collapse;display:block;overflow:auto}th,td{border:1px solid #c3d0d8;padding:10px;text-align:left}nav,.toc{background:#f3f7fa;padding:16px;margin-bottom:24px}nav{display:flex;gap:18px;flex-wrap:wrap}img{max-width:100%}a:focus-visible{outline:3px solid #155cb0;outline-offset:3px}@media print{nav,.toc{display:none}body{padding:0;max-width:none}h2,h3{break-after:avoid}}'''
NAV = '''<nav aria-label="Documentation"><a href="/">ScriptStudio</a><a href="/user-manual.html">User manual</a><a href="/ScriptStudio-User-Manual.docx" download>Download Word manual (.docx)</a><a href="/README.html">README</a><a href="/README.md" download>Download README.md</a></nav>'''


def render(source, title, target):
    converter = markdown.Markdown(extensions=['extra', 'toc'])
    body = converter.convert(source.read_text())
    tree = html.fragment_fromstring(body, create_parent='main')
    for element in tree.xpath('.//*[@href] | .//*[@src]'):
        attribute = 'href' if element.get('href') is not None else 'src'
        url = element.get(attribute)
        if url.startswith(('#', '/')) or urlsplit(url).scheme:
            continue
        if url == 'docs/user-manual.md':
            url = '/user-manual.html'
        elif url.startswith('frontend/public/'):
            url = '/' + url.removeprefix('frontend/public/')
        else:
            url = SOURCE_URL + url
        element.set(attribute, url)
    page = f'<!doctype html>\n<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{escape(title)}</title><style>{CSS}</style></head><body>{NAV}<aside aria-label="Contents">{converter.toc}</aside>{html.tostring(tree, encoding="unicode")}\n</body></html>\n'
    (PUBLIC / target).write_text(page)
    return tree


def main():
    PUBLIC.mkdir(parents=True, exist_ok=True)
    manual = render(ROOT / 'docs/user-manual.md', 'ScriptStudio User Manual', 'user-manual.html')
    render(ROOT / 'README.md', 'ScriptStudio README', 'README.html')
    shutil.copyfile(ROOT / 'README.md', PUBLIC / 'README.md')
    document = Document()
    document.core_properties.title = 'ScriptStudio User Manual'
    document.core_properties.author = 'ScriptStudio'
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    for element in manual:
        text = element.text_content()
        if element.tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            document.add_heading(text, level=int(element.tag[1]))
        elif element.tag in ('ul', 'ol'):
            for item in element:
                document.add_paragraph(item.text_content(), style='List Bullet' if element.tag == 'ul' else 'List Number')
        elif element.tag == 'p':
            document.add_paragraph(text)
        else:
            raise ValueError(f'Unsupported manual element: {element.tag}')
    document.save(PUBLIC / 'ScriptStudio-User-Manual.docx')
    print('Generated browser manual, Word manual, browser README, and downloadable README.')


if __name__ == '__main__':
    main()
